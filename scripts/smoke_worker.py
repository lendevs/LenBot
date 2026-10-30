"""Run one task through a real task sandbox: Pi, the model bridge, a question and a delivery.

The model is a loopback Chat Completions endpoint that answers from a fixed
script: run a shell command, ask the requester, deliver the file, then finish.
Everything else is the task service, the sandbox and Pi as deployed: a worker
image under Docker, or a task runtime (``build_task_runtime.py``) under bwrap.
``--network-probe`` adds curl and Node fetch steps; ``--egress-mode`` picks the mode they run in
(allowlist with only example.com listed by default, public, or off).
``--copy-through-docker`` reads the delivery with ``docker cp`` as on Windows.
``--workloads`` adds a step that handles CSV/XLSX, draws a Chinese chart, writes DOCX and, where
``lenbot-render`` exists, turns HTML into PDF, all without network access.
``--task-volume NAME:ROOT`` runs inside a host container whose instance volume NAME is mounted at ROOT;
task containers then mount their directories as subpaths of that volume.
"""

import argparse
import asyncio
import json
import os
from pathlib import Path
import platform
import shutil
import subprocess
import sys
import tempfile

from len_bot.next.config import HostConfig
from len_bot.next.storage.store import Store
from len_bot.next.work.sandbox import DockerSandbox
from len_bot.next.work.service import WorkTasks
from len_bot.next.work.store import TaskStore

SCENE = 'onebot:group:80001'
REQUESTER = 'onebot:70001'
REPORT = 'worker smoke report'
PROBE = ("for target in https://example.com/ http://not-listed.invalid/ http://10.0.0.1/; do "
         "printf '%s=%s\\n' \"$target\" \"$(curl -s -o /dev/null -m 20 -w '%{http_code}/%{http_connect}' \"$target\")\"; done\n"
         # Node's fetch() reads the proxy variables from the environment of the task.
         "for target in https://example.com/ http://not-listed.invalid/; do "
         "printf 'node-fetch %s=%s\\n' \"$target\" \"$(node -e \"fetch(process.argv[1], {signal: AbortSignal.timeout(20000)})"
         ".then(response => console.log(response.status), () => console.log('error'))\" \"$target\")\"; done")
EXPECTED_PROBE = {
    'allowlist': ('https://example.com/=200/200', 'http://not-listed.invalid/=502/000', 'http://10.0.0.1/=502/000',
                  'node-fetch https://example.com/=200', 'node-fetch http://not-listed.invalid/=error'),
    'public': ('https://example.com/=200/200', 'http://10.0.0.1/=502/000',
               'node-fetch https://example.com/=200', 'node-fetch http://not-listed.invalid/=error'),
    'off': ('https://example.com/=000/000', 'http://not-listed.invalid/=000/000', 'http://10.0.0.1/=000/000',
            'node-fetch https://example.com/=error', 'node-fetch http://not-listed.invalid/=error'),
}
WORKLOADS = r'''python3 - <<'PY' 2>&1
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import openpyxl
import pandas as pd
from docx import Document
from PIL import Image
pd.DataFrame({'城市': ['北京', '上海'], '销量': [3, 5]}).to_csv('out/data.csv', index=False)
frame = pd.read_csv('out/data.csv')
frame.to_excel('out/data.xlsx', index=False)
assert openpyxl.load_workbook('out/data.xlsx').active['A2'].value == '北京'
plt.bar(frame['城市'], frame['销量'])
plt.title('各城市销量')
plt.savefig('out/chart.png')
document = Document()
document.add_heading('销量报告')
document.add_picture('out/chart.png')
document.save('out/report.docx')
with Image.open('out/chart.png') as chart:
    print('workloads', int(frame['销量'].sum()), chart.size)
PY
if command -v lenbot-render >/dev/null; then
  printf '<meta charset="utf-8"><h1>销量报告</h1><p>北京 3，上海 5</p>' > page.html
  lenbot-render page.html --pdf out/page.pdf >/dev/null
  python3 -c "from pypdf import PdfReader; print('pdf pages', len(PdfReader('out/page.pdf').pages))"
fi'''


def chunk(body: dict) -> bytes:
    return b'data: ' + json.dumps(body, ensure_ascii=False).encode() + b'\n\n'


def scripted_reply(request: dict, outputs: list[str], probe: bool, workloads: bool) -> bytes:
    """One streamed step chosen by how many tool results the conversation already has."""
    tools = {tool['function']['name']: tool['function'] for tool in request.get('tools', [])}
    results = [message for message in request['messages'] if message['role'] == 'tool']
    steps = [
        ('bash', {'command': f"mkdir -p out && printf '%s\\n' '{REPORT}' \"uid=$(id -u)\" > out/report.txt"}),
        ('ask_requester', {'question': '报告用什么标题？'}),
        ('deliver_file', {'path': 'out/report.txt', 'name': '报告.txt', 'note': 'worker 冒烟'}),
        *([('bash', {'command': PROBE})] if probe else []),
        *([('bash', {'command': WORKLOADS})] if workloads else []),
    ]
    usage = {'choices': [], 'usage': {'prompt_tokens': 50, 'completion_tokens': 5}}
    outputs[:] = [result['content'] if isinstance(result['content'], str)
                  else json.dumps(result['content'], ensure_ascii=False) for result in results]
    if len(results) < len(steps):
        name, arguments = steps[len(results)]
        missing = set(arguments) - set(tools[name]['parameters']['properties'])
        if missing:
            raise ValueError(f'Pi tool {name} lacks parameters {sorted(missing)}: {tools[name]}')
        call = {'index': 0, 'id': f'call_{len(results) + 1}', 'type': 'function',
                'function': {'name': name, 'arguments': json.dumps(arguments, ensure_ascii=False)}}
        body = [chunk({'choices': [{'index': 0, 'delta': {'role': 'assistant', 'tool_calls': [call]},
                                    'finish_reason': None}]}),
                chunk({'choices': [{'index': 0, 'delta': {}, 'finish_reason': 'tool_calls'}]})]
    else:
        body = [chunk({'choices': [{'index': 0, 'delta': {'role': 'assistant', 'content': '完成'},
                                    'finish_reason': 'stop'}]})]
    return b''.join(body) + chunk(usage) + b'data: [DONE]\n\n'


async def serve_model(requests: list[dict], outputs: list[str], probe: bool, workloads: bool):
    async def respond(reader, writer):
        header = (await reader.readuntil(b'\r\n\r\n')).decode()
        fields = {name.lower(): value for name, _, value in
                  (line.partition(': ') for line in header.split('\r\n')[1:]) if value}
        request = json.loads(await reader.readexactly(int(fields['content-length'])))
        requests.append(request)
        try:
            stream, status = scripted_reply(request, outputs, probe, workloads), b'200 OK'
        except Exception as error:
            stream, status = json.dumps({'error': {'message': str(error)}}).encode(), b'400 Bad Request'
        writer.write(b'HTTP/1.1 ' + status + b'\r\nContent-Type: text/event-stream\r\n'
                     + f'Content-Length: {len(stream)}\r\nConnection: close\r\n\r\n'.encode() + stream)
        await writer.drain()
        writer.close()
    return await asyncio.start_server(respond, '127.0.0.1', 0)


def docker_host() -> str:
    configured = os.environ.get('DOCKER_HOST', '')
    if configured.startswith('unix:///'):
        return configured
    desktop = Path.home() / '.docker/run/docker.sock'
    return f'unix://{desktop}' if desktop.exists() else 'unix:///var/run/docker.sock'


def task_containers(docker: str, host: str) -> set[str]:
    output = subprocess.run([docker, '--host', host, 'ps', '-aq', '--no-trunc', '--filter', 'name=lenbot-next-'],
                            check=True, capture_output=True, text=True).stdout
    return set(output.split())


async def until(condition, seconds: float):
    async with asyncio.timeout(seconds):
        while not (result := condition()):
            await asyncio.sleep(.1)
    return result


async def smoke(sandbox: dict, root: Path, network: dict | None = None, volume: dict | None = None,
                workloads: bool = False, mode: str = 'allowlist') -> dict:
    """``network`` turns on the probe; it holds this machine's network settings, empty when defaults fit."""
    requests: list[dict] = []
    outputs: list[str] = []
    probe = network is not None
    server = await serve_model(requests, outputs, probe, workloads)
    port = server.sockets[0].getsockname()[1]
    config = HostConfig.model_validate({
        'compaction': {'input_tokens': 2000}, 'mode': 'isolated-multi', 'bot_id': 'onebot:90001', 'timezone': 'UTC',
        'owners': [REQUESTER], 'database': root / 'host.sqlite3', 'delivery': 'simulated', 'onebot': None,
        **({'network': network} if network else {}),
        'models': {
            'providers': {'smoke': {'api': 'openai-chat', 'base_url': f'http://127.0.0.1:{port}/v1',
                                    'api_key': 'synthetic-upstream-key'}},
            'roles': {'mind': {'provider': 'smoke', 'model': 'scripted-worker', 'context_window_tokens': 65536}},
        },
        'worker': {
            **sandbox, 'model_reasoning': False,
            'workspace_root': root / 'workspaces', 'runtime_root': root / 'runtime',
            'delivery_root': root / 'deliveries',
            'egress': ({'mode': 'allowlist', 'allowed_hosts': ['example.com']} if mode == 'allowlist'
                       else {'mode': mode}) if probe else {'mode': 'off'},
            **({'task_volume': volume} if volume else {}),
        },
        'scenes': {SCENE: {'persona': root / 'unused-role', 'tasks': {'enabled': True}}},
    })
    docker = sandbox.get('sandbox', 'docker') == 'docker'
    before = task_containers(str(sandbox['docker_binary']), sandbox['docker_host']) if docker else set()
    async with server:
        with Store(config.database) as store:
            tasks = WorkTasks(config, store, None, lambda scene: None, skills={SCENE: ()}, memory=None,
                              data_tools={SCENE: []}, skill_permissions={SCENE: []}, tool_permissions={SCENE: []})
            await tasks.recover()
            await tasks.start()
            try:
                created = await tasks.delegate(SCENE, requester=REQUESTER, goal='写一份冒烟报告',
                                               deliverable='报告.txt', context='worker 镜像冒烟')
                id = created['id']
                status = lambda: tasks.status(SCENE, id)
                waiting = await until(lambda: status()['status'] in {'waiting_input', 'done', 'failed'} and status(), 180)
                if waiting['status'] != 'waiting_input':
                    raise RuntimeError(f'task did not ask its question: {json.dumps(waiting, ensure_ascii=False)}')
                question = waiting['question']
                await tasks.answer(SCENE, id, requester=REQUESTER, text='冒烟标题', confirmed=None,
                                   question_id=question['id'])
                done = await until(lambda: status()['status'] in {'done', 'failed', 'cancelled'} and status(), 180)
                files = TaskStore(store).list_files(SCENE, id)
            finally:
                await tasks.close()
    after = task_containers(str(sandbox['docker_binary']), sandbox['docker_host']) if docker else set()
    if done['status'] != 'done':
        raise RuntimeError(f'task ended {done["status"]}: {done["error"]}')
    [file] = files
    report = Path(file.path).read_text(encoding='utf-8')
    checks = {
        'question': (question['method'], question['title']) == ('input', '报告用什么标题？'),
        'answer_reached_model': len(outputs) > 1 and '冒烟标题' in outputs[1],
        'delivery': report.startswith(REPORT) and f"uid={sandbox['uid'] if docker else 10000}" in report,
        'model_calls': done['model_calls'] == 4 + probe + workloads and len(requests) == 4 + probe + workloads,
        **({'network': len(outputs) > 3 and all(line in outputs[3] for line in EXPECTED_PROBE[mode])} if probe else {}),
        **({'workloads': len(outputs) > 3 + probe and 'workloads 8 ' in outputs[3 + probe]
            # The local sandbox uses the system's fonts; the worker image ships Noto CJK.
            and ('missing from' not in outputs[3 + probe] or not docker)
            and ('pdf pages 1' in outputs[3 + probe] or not docker)} if workloads else {}),
        'sandbox_removed': after <= before and done['container'] is None,
    }
    if not all(checks.values()):
        raise RuntimeError(f'worker smoke checks failed: {checks}; tool outputs={outputs[3:]!r}; '
                           f'status={json.dumps(done, ensure_ascii=False)}')
    result = {'host': f'{platform.system()} {platform.machine()}', 'summary': done['summary'], 'checks': checks,
              **({'network_probe': outputs[3]} if probe else {}),
              **({'workloads_output': outputs[3 + probe]} if workloads else {})}
    if docker:
        cli = [str(sandbox['docker_binary']), '--host', sandbox['docker_host']]
        inspect = json.loads(subprocess.run([*cli, 'image', 'inspect', sandbox['image']],
                                            check=True, capture_output=True, text=True).stdout)[0]
        version = subprocess.run([*cli, 'run', '--rm', '--network', 'none', '--entrypoint', 'pi', sandbox['image'],
                                  '--version'], check=True, capture_output=True, text=True).stdout.strip()
        return {**result, 'image': sandbox['image'], 'image_id': inspect['Id'],
                'architecture': inspect['Architecture'], 'pi': version}
    runtime = sandbox['task_runtime'] / 'node' / 'bin'
    version = subprocess.run([str(runtime / 'node'), str(runtime / 'pi'), '--version'], check=True,
                             capture_output=True, text=True).stdout.strip()
    bwrap = subprocess.run([str(sandbox['bwrap_binary']), '--version'], check=True, capture_output=True,
                           text=True).stdout.strip()
    return {**result, 'task_runtime': str(sandbox['task_runtime']), 'pi': version, 'bwrap': bwrap}


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--sandbox', choices=['docker', 'bwrap'], default='docker')
    parser.add_argument('--image', default='lenbot-worker:local')
    parser.add_argument('--docker', default=shutil.which('docker'))
    parser.add_argument('--docker-host', default=docker_host())
    parser.add_argument('--bwrap', default=shutil.which('bwrap'))
    parser.add_argument('--task-runtime', type=Path, help='directory built by build_task_runtime.py')
    parser.add_argument('--network-probe', nargs='?', const='{}', metavar='NETWORK_JSON',
                        help="add the curl and Node fetch steps; optional JSON of this machine's network settings")
    parser.add_argument('--egress-mode', choices=sorted(EXPECTED_PROBE), default='allowlist',
                        help='the egress mode the --network-probe steps run in')
    parser.add_argument('--copy-through-docker', action='store_true', help='read the delivery as on Windows')
    parser.add_argument('--task-volume', metavar='NAME:ROOT', help='instance volume and its mount in this container')
    parser.add_argument('--workloads', action='store_true', help='add the offline file-handling step')
    arguments = parser.parse_args()
    if arguments.copy_through_docker:
        DockerSandbox.copy_through_docker = True
    if arguments.sandbox == 'docker':
        if arguments.docker is None:
            sys.exit('docker CLI not found; pass --docker')
        sandbox = {'docker_binary': Path(arguments.docker), 'docker_host': arguments.docker_host,
                   'image': arguments.image, 'uid': os.getuid(), 'gid': os.getgid()}
    else:
        if arguments.bwrap is None or arguments.task_runtime is None:
            sys.exit('bwrap mode needs the bwrap binary and --task-runtime')
        sandbox = {'sandbox': 'bwrap', 'bwrap_binary': Path(arguments.bwrap),
                   'task_runtime': arguments.task_runtime.resolve()}
    volume = None
    if arguments.task_volume:
        name, _, mount = arguments.task_volume.partition(':')
        volume = {'name': name, 'root': Path(mount).resolve()}
    with tempfile.TemporaryDirectory(prefix='lenbot-worker-smoke-', dir=volume and volume['root']) as directory:
        result = asyncio.run(smoke(sandbox, Path(directory).resolve(),
                                   None if arguments.network_probe is None else json.loads(arguments.network_probe),
                                   volume, arguments.workloads, arguments.egress_mode))
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
